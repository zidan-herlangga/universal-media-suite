import os
import platform
import shutil
import subprocess
import sys
import threading
from collections import Counter
from io import BytesIO
from tkinter import filedialog, messagebox
import customtkinter as ctk

# Integrasi Drag and Drop Tkinter
from tkinterdnd2 import DND_FILES, TkinterDnD

# -------------------------------------------------------------
# SAFE IMPORT BACKEND ENGINE
# -------------------------------------------------------------
try:
  from PIL import Image
  import pillow_heif
  import rawpy

  pillow_heif.register_heif_opener()
  HAS_IMAGE_ENGINE = True
except ImportError:
  HAS_IMAGE_ENGINE = False

# COM Windows hanya tersedia di OS Windows
try:
  import pythoncom
  import win32com.client

  HAS_WIN32COM = True
except ImportError:
  HAS_WIN32COM = False

try:
  from pdf2docx import Converter as PdfToDocxConverter

  HAS_PDF2DOCX = True
except ImportError:
  HAS_PDF2DOCX = False

try:
  import pymupdf as fitz  # PyMuPDF modern (render halaman PDF jadi gambar)

  HAS_FITZ = True
except ImportError:
  try:
    import fitz  # PyMuPDF lama

    HAS_FITZ = True
  except ImportError:
    HAS_FITZ = False

try:
  from docx import Document as DocxDocument
  from docx.enum.section import WD_SECTION
  from docx.shared import Mm

  HAS_DOCX = True
except ImportError:
  HAS_DOCX = False

try:
  from pypdf import PdfReader, PdfWriter

  HAS_PYPDF = True
except ImportError:
  HAS_PYPDF = False

try:
  import yt_dlp

  HAS_YTDLP = True
except ImportError:
  HAS_YTDLP = False

# Konfigurasi Tema Global
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")


# -------------------------------------------------------------
# HELPER CROSS-PLATFORM
# -------------------------------------------------------------
def get_ffmpeg_binary():
  """Mengembalikan lokasi binary ffmpeg: bundle PyInstaller / folder aplikasi / PATH OS."""
  bundle_dir = getattr(sys, "_MEIPASS", None)
  for base in (bundle_dir, os.path.dirname(os.path.abspath(sys.executable))):
    if not base:
      continue
    for root in (base, os.path.join(base, "_internal")):
      for name in ("ffmpeg.exe", "ffmpeg"):
        p = os.path.join(root, name)
        if os.path.exists(p):
          return p
  return shutil.which("ffmpeg")


def check_ffmpeg_available():
  """Mengecek binary ffmpeg tersedia di bundle lokal maupun PATH OS (Windows/Linux/Mac)."""
  return get_ffmpeg_binary() is not None


def parse_dnd_paths(event_data):
  """Mem-parse string path dari TkinterDnD secara aman lintas OS."""
  raw_paths = event_data.strip()
  if "{" in raw_paths:
    import re

    paths = re.findall(r"\{([^}]+)\}|(\S+)", raw_paths)
    paths = [p[0] if p[0] else p[1] for p in paths]
  else:
    paths = raw_paths.split()
  return [os.path.normpath(p) for p in paths]


def open_folder_in_explorer(folder_path):
  """Membuka folder di File Explorer (Windows), Finder (macOS), atau File Manager (Linux)."""
  if not os.path.exists(folder_path):
    return
  current_os = platform.system()
  try:
    if current_os == "Windows":
      os.startfile(folder_path)
    elif current_os == "Darwin":  # macOS
      subprocess.run(["open", folder_path])
    else:  # Linux
      subprocess.run(["xdg-open", folder_path])
  except Exception as e:
    print(f"Gagal membuka folder: {e}")


def get_libreoffice_command():
  """Mencari binary LibreOffice untuk Linux, macOS, atau Windows non-MS Office."""
  candidates = [
      "soffice",
      "libreoffice",
      "/Applications/LibreOffice.app/Contents/MacOS/soffice",
      r"C:\Program Files\LibreOffice\program\soffice.exe",
      r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
  ]
  for cmd in candidates:
    if shutil.which(cmd) or os.path.exists(cmd):
      return cmd
  return None


# =============================================================
# TAB 1: MODUL FOTO & RAW (DND SUPPORTED)
# =============================================================
class ImageTab(ctk.CTkFrame):

  def __init__(self, parent):
    super().__init__(parent, fg_color="transparent")
    self.source_label_text = ctk.StringVar(
        value="Seret foto/folder ke sini atau klik tombol pilih..."
    )
    self.folder_output = ctk.StringVar()
    self.target_format = ctk.StringVar(value="JPG")

    self.enable_resize = ctk.BooleanVar(value=False)
    self.width_var = ctk.StringVar(value="1920")
    self.height_var = ctk.StringVar(value="1080")
    self.keep_aspect = ctk.BooleanVar(value=True)

    self.cancel_event = threading.Event()
    self.selected_file_paths = []

    self.raw_exts = {
        "cr2", "cr3", "crw", "nef", "nrw", "arw", "srf", "sr2", "raf",
        "orf", "rw2", "pef", "ptx", "dng", "x3f", "erf", "mef", "mos",
        "kdc", "dcr", "raw",
    }
    self.pillow_exts = {
        "jpg", "jpeg", "png", "webp", "bmp", "tiff", "tif", "gif",
        "ico", "tga", "ppm", "pcx", "heic", "heif", "hif", "psd",
    }
    self.supported_exts = self.raw_exts.union(self.pillow_exts)

    self.setup_ui()

  def setup_ui(self):
    self.drop_zone = ctk.CTkFrame(
        self, corner_radius=12, border_width=2, border_color="#334155"
    )
    self.drop_zone.pack(fill="x", padx=10, pady=(10, 8))

    ctk.CTkLabel(
        self.drop_zone,
        text="SUMBER FOTO / RAW (DRAG & DROP ATAU PILIH)",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=14, pady=(10, 2))

    f_in = ctk.CTkFrame(self.drop_zone, fg_color="transparent")
    f_in.pack(fill="x", padx=14, pady=(0, 6))

    self.entry_source = ctk.CTkEntry(
        f_in, textvariable=self.source_label_text, state="readonly", height=34
    )
    self.entry_source.pack(side="left", fill="x", expand=True, padx=(0, 8))

    ctk.CTkButton(
        f_in,
        text="File...",
        width=70,
        height=34,
        command=self.select_files,
        cursor="hand2",
    ).pack(side="left", padx=(0, 6))
    ctk.CTkButton(
        f_in,
        text="Folder...",
        width=75,
        height=34,
        fg_color="#334155",
        hover_color="#475569",
        command=self.select_folder,
        cursor="hand2",
    ).pack(side="right")

    self.lbl_detect = ctk.CTkLabel(
        self.drop_zone,
        text="Tarik file foto satuan, banyak foto, atau 1 folder penuh ke sini.",
        text_color="#64748b",
        font=ctk.CTkFont(size=11),
    )
    self.lbl_detect.pack(anchor="w", padx=14, pady=(0, 6))

    ctk.CTkLabel(
        self.drop_zone,
        text="FOLDER SIMPAN HASIL",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=14, pady=(4, 2))
    f_out = ctk.CTkFrame(self.drop_zone, fg_color="transparent")
    f_out.pack(fill="x", padx=14, pady=(0, 12))
    ctk.CTkEntry(
        f_out, textvariable=self.folder_output, state="readonly", height=34
    ).pack(side="left", fill="x", expand=True, padx=(0, 8))
    ctk.CTkButton(
        f_out,
        text="Pilih",
        width=75,
        height=34,
        command=self.select_out,
        cursor="hand2",
    ).pack(side="right")

    self.drop_zone.drop_target_register(DND_FILES)
    self.drop_zone.dnd_bind("<<Drop>>", self.on_drop)
    self.entry_source.drop_target_register(DND_FILES)
    self.entry_source.dnd_bind("<<Drop>>", self.on_drop)

    card_opt = ctk.CTkFrame(self, corner_radius=12)
    card_opt.pack(fill="x", padx=10, pady=(0, 10))

    f_opt = ctk.CTkFrame(card_opt, fg_color="transparent")
    f_opt.pack(fill="x", padx=14, pady=10)
    ctk.CTkLabel(
        f_opt, text="Target Format:", font=ctk.CTkFont(size=12, weight="bold")
    ).pack(side="left", padx=(0, 8))
    ctk.CTkOptionMenu(
        f_opt,
        variable=self.target_format,
        values=["JPG", "PNG", "WEBP", "PDF", "TIFF", "BMP", "ICO", "TGA"],
        width=110,
    ).pack(side="left")

    ctk.CTkSwitch(
        card_opt,
        text="Aktifkan Resize Dimensi Maksimal",
        variable=self.enable_resize,
        command=self.toggle_resize,
    ).pack(anchor="w", padx=14, pady=(0, 6))

    f_dim = ctk.CTkFrame(card_opt, fg_color="transparent")
    f_dim.pack(fill="x", padx=14, pady=(0, 12))
    ctk.CTkLabel(f_dim, text="Lebar:").pack(side="left", padx=(0, 4))
    self.e_w = ctk.CTkEntry(
        f_dim, textvariable=self.width_var, width=60, state="disabled"
    )
    self.e_w.pack(side="left", padx=(0, 10))
    ctk.CTkLabel(f_dim, text="Tinggi:").pack(side="left", padx=(0, 4))
    self.e_h = ctk.CTkEntry(
        f_dim, textvariable=self.height_var, width=60, state="disabled"
    )
    self.e_h.pack(side="left", padx=(0, 10))
    self.c_asp = ctk.CTkCheckBox(
        f_dim,
        text="Jaga Aspek Rasio",
        variable=self.keep_aspect,
        state="disabled",
    )
    self.c_asp.pack(side="left")

    self.lbl_status = ctk.CTkLabel(
        self, text="Siap.", text_color="#94a3b8", font=ctk.CTkFont(size=12)
    )
    self.lbl_status.pack(anchor="w", padx=14, pady=(2, 2))
    self.progress = ctk.CTkProgressBar(self, height=8, corner_radius=4)
    self.progress.pack(fill="x", padx=10, pady=(0, 12))
    self.progress.set(0)

    f_btn = ctk.CTkFrame(self, fg_color="transparent")
    f_btn.pack(fill="x", padx=10, pady=(0, 10))
    self.btn_run = ctk.CTkButton(
        f_btn,
        text="Mulai Konversi Foto",
        height=40,
        font=ctk.CTkFont(size=13, weight="bold"),
        state="disabled",
        cursor="hand2",
        command=self.start_batch,
    )
    self.btn_run.pack(side="left", fill="x", expand=True, padx=(0, 6))
    self.btn_stop = ctk.CTkButton(
        f_btn,
        text="Hentikan",
        height=40,
        font=ctk.CTkFont(size=13, weight="bold"),
        fg_color="#dc2626",
        hover_color="#b91c1c",
        state="disabled",
        cursor="hand2",
        command=self.stop_batch,
    )
    self.btn_stop.pack(side="right", padx=(6, 0))

  def on_drop(self, event):
    dropped_paths = parse_dnd_paths(event.data)
    valid_files = []
    counter = Counter()

    for item in dropped_paths:
      if os.path.isdir(item):
        for root, _, files in os.walk(item):
          for f in files:
            ext = os.path.splitext(f)[1].lower().replace(".", "")
            if ext in self.supported_exts:
              valid_files.append(os.path.join(root, f))
              counter[ext] += 1
      elif os.path.isfile(item):
        ext = os.path.splitext(item)[1].lower().replace(".", "")
        if ext in self.supported_exts:
          valid_files.append(item)
          counter[ext] += 1

    if valid_files:
      self.selected_file_paths = valid_files
      self.source_label_text.set(
          valid_files[0]
          if len(valid_files) == 1
          else f"{len(valid_files)} file foto terpilih"
      )
      if not self.folder_output.get():
        self.folder_output.set(os.path.dirname(valid_files[0]))

      info = ", ".join([f"{k.upper()}:{v}" for k, v in counter.items()])
      self.lbl_detect.configure(
          text=f"Terdeteksi {len(valid_files)} foto ({info})",
          text_color="#38bdf8",
      )
      self.check_state()
    else:
      self.lbl_detect.configure(
          text="File yang di-drop bukan format foto yang didukung.",
          text_color="#ef4444",
      )

  def toggle_resize(self):
    st = "normal" if self.enable_resize.get() else "disabled"
    self.e_w.configure(state=st)
    self.e_h.configure(state=st)
    self.c_asp.configure(state=st)

  def select_files(self):
    all_ext_filter = " ".join([f"*.{e}" for e in sorted(self.supported_exts)])
    files = filedialog.askopenfilenames(
        title="Pilih File Foto / RAW (Bisa Banyak)",
        filetypes=[
            ("Semua Foto & RAW", all_ext_filter),
            ("Semua File", "*.*"),
        ],
    )
    if files:
      self.selected_file_paths = list(files)
      self.source_label_text.set(
          files[0] if len(files) == 1 else f"{len(files)} file foto terpilih"
      )
      if not self.folder_output.get():
        self.folder_output.set(os.path.dirname(files[0]))
      self.lbl_detect.configure(
          text=f"Terpilih {len(files)} foto.", text_color="#38bdf8"
      )
      self.check_state()

  def select_folder(self):
    d = filedialog.askdirectory(title="Pilih Folder Sumber Foto")
    if d:
      self.source_label_text.set(f"Folder: {d}")
      if not self.folder_output.get():
        self.folder_output.set(d)
      files = [
          os.path.join(d, f)
          for f in os.listdir(d)
          if os.path.isfile(os.path.join(d, f))
          and os.path.splitext(f)[1].lower().replace(".", "")
          in self.supported_exts
      ]
      self.selected_file_paths = files
      self.lbl_detect.configure(
          text=f"Ditemukan {len(files)} foto.",
          text_color="#38bdf8" if files else "#ef4444",
      )
      self.check_state()

  def select_out(self):
    d = filedialog.askdirectory(title="Pilih Folder Tujuan Simpan")
    if d:
      self.folder_output.set(d)
      self.check_state()

  def check_state(self):
    if self.selected_file_paths and self.folder_output.get().strip():
      self.btn_run.configure(state="normal")
      self.lbl_status.configure(
          text=f"Siap memproses {len(self.selected_file_paths)} file.",
          text_color="#f8fafc",
      )
    else:
      self.btn_run.configure(state="disabled")

  def stop_batch(self):
    self.cancel_event.set()
    self.btn_stop.configure(state="disabled")

  def start_batch(self):
    self.cancel_event.clear()
    self.btn_run.configure(state="disabled")
    self.btn_stop.configure(state="normal")
    threading.Thread(target=self._process, daemon=True).start()

  def _process(self):
    out_dir = self.folder_output.get()
    t_ext = self.target_format.get().lower()
    save_fmt = "JPEG" if t_ext in ["jpg", "jpeg"] else t_ext.upper()

    total = len(self.selected_file_paths)
    ok_count = 0
    failed_items = []
    cancelled = False

    for idx, src_p in enumerate(self.selected_file_paths, start=1):
      if self.cancel_event.is_set():
        cancelled = True
        break

      fname = os.path.basename(src_p)
      ext = os.path.splitext(fname)[1].lower().replace(".", "")
      self.lbl_status.configure(text=f"Memproses ({idx}/{total}): {fname}")
      dest_p = os.path.join(out_dir, f"{os.path.splitext(fname)[0]}.{t_ext}")

      try:
        if ext in self.raw_exts:
          with rawpy.imread(src_p) as raw:
            rgb = raw.postprocess(use_camera_wb=True)
          img = Image.fromarray(rgb)
        else:
          img = Image.open(src_p)

        if self.enable_resize.get():
          tw, th = int(self.width_var.get()), int(self.height_var.get())
          if self.keep_aspect.get():
            img.thumbnail((tw, th), Image.Resampling.LANCZOS)
          else:
            img = img.resize((tw, th), Image.Resampling.LANCZOS)

        if t_ext in ["jpg", "jpeg", "pdf", "bmp"] and img.mode in (
            "RGBA", "LA", "P",
        ):
          img = img.convert("RGB")
        elif t_ext == "ico" and max(img.size) > 256:
          img.thumbnail((256, 256), Image.Resampling.LANCZOS)

        img.save(dest_p, format=save_fmt, quality=95)
        ok_count += 1
      except Exception as ex:
        failed_items.append((fname, str(ex)))

      self.progress.set(idx / total)

    self.btn_run.configure(state="normal")
    self.btn_stop.configure(state="disabled")

    if failed_items:
      log_path = os.path.join(out_dir, "conversion_error_log.txt")
      with open(log_path, "w", encoding="utf-8") as lf:
        lf.write("LAPORAN FILE GAGAL DIKONVERSI:\n" + "=" * 40 + "\n")
        for fn, err in failed_items:
          lf.write(f"- {fn}: {err}\n")

    msg = f"Berhasil memproses {ok_count} dari {total} file."
    if failed_items:
      msg += f"\n\nPerhatian: {len(failed_items)} file gagal dikonversi (dicatat di conversion_error_log.txt)."

    msg += "\n\nBuka folder hasil sekarang?"
    if messagebox.askyesno("Hasil Konversi Foto", msg):
      open_folder_in_explorer(out_dir)


# =============================================================
# TAB 2: MODUL VIDEO RESIZER (DND SUPPORTED)
# =============================================================
class VideoTab(ctk.CTkFrame):

  def __init__(self, parent, ffmpeg_available):
    super().__init__(parent, fg_color="transparent")
    self.ffmpeg_available = ffmpeg_available
    self.source_label_text = ctk.StringVar(
        value="Seret file video atau folder ke sini..."
    )
    self.folder_output = ctk.StringVar()
    self.preset_var = ctk.StringVar(value="1280x720 (720p HD)")
    self.custom_w = ctk.StringVar(value="1280")
    self.custom_h = ctk.StringVar(value="720")

    self.current_process = None
    self.cancel_event = threading.Event()
    self.valid_files = []
    self.video_exts = {
        "mp4", "mkv", "mov", "avi", "webm", "flv", "wmv", "m4v", "ts", "3gp", "mpg",
    }

    self.setup_ui()

  def setup_ui(self):
    if not self.ffmpeg_available:
      f_warn = ctk.CTkFrame(self, fg_color="#ef4444", corner_radius=8)
      f_warn.pack(fill="x", padx=10, pady=(6, 0))
      ctk.CTkLabel(
          f_warn,
          text=(
              "⚠️ Binary FFmpeg tidak ditemukan! Pastikan ffmpeg terinstal di sistem Anda."
          ),
          text_color="white",
          font=ctk.CTkFont(size=11, weight="bold"),
      ).pack(pady=6)

    self.drop_zone = ctk.CTkFrame(
        self, corner_radius=12, border_width=2, border_color="#334155"
    )
    self.drop_zone.pack(fill="x", padx=10, pady=(10, 8))

    ctk.CTkLabel(
        self.drop_zone,
        text="SUMBER VIDEO (DRAG & DROP / PILIH)",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=14, pady=(10, 2))

    f_in = ctk.CTkFrame(self.drop_zone, fg_color="transparent")
    f_in.pack(fill="x", padx=14, pady=(0, 6))
    self.entry_source = ctk.CTkEntry(
        f_in, textvariable=self.source_label_text, state="readonly", height=34
    )
    self.entry_source.pack(side="left", fill="x", expand=True, padx=(0, 8))
    ctk.CTkButton(
        f_in,
        text="Pilih",
        width=75,
        height=34,
        command=self.select_files,
        cursor="hand2",
    ).pack(side="right")

    self.lbl_detect = ctk.CTkLabel(
        self.drop_zone,
        text="Tarik file video atau folder ke sini.",
        text_color="#64748b",
        font=ctk.CTkFont(size=11),
    )
    self.lbl_detect.pack(anchor="w", padx=14, pady=(0, 6))

    ctk.CTkLabel(
        self.drop_zone,
        text="FOLDER TUJUAN SIMPAN",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=14, pady=(4, 2))
    f_out = ctk.CTkFrame(self.drop_zone, fg_color="transparent")
    f_out.pack(fill="x", padx=14, pady=(0, 12))
    ctk.CTkEntry(
        f_out, textvariable=self.folder_output, state="readonly", height=34
    ).pack(side="left", fill="x", expand=True, padx=(0, 8))
    ctk.CTkButton(
        f_out,
        text="Pilih",
        width=75,
        height=34,
        command=self.select_out,
        cursor="hand2",
    ).pack(side="right")

    self.drop_zone.drop_target_register(DND_FILES)
    self.drop_zone.dnd_bind("<<Drop>>", self.on_drop)
    self.entry_source.drop_target_register(DND_FILES)
    self.entry_source.dnd_bind("<<Drop>>", self.on_drop)

    card_res = ctk.CTkFrame(self, corner_radius=12)
    card_res.pack(fill="x", padx=10, pady=(0, 10))

    ctk.CTkLabel(
        card_res,
        text="TARGET RESOLUSI VIDEO",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=14, pady=(10, 4))
    presets = [
        "1920x1080 (1080p FHD)",
        "1280x720 (720p HD)",
        "854x480 (480p SD)",
        "640x360 (360p)",
        "Custom (Manual)",
    ]
    ctk.CTkOptionMenu(
        card_res,
        variable=self.preset_var,
        values=presets,
        command=self.on_preset_change,
        height=34,
    ).pack(fill="x", padx=14, pady=(0, 10))

    f_dim = ctk.CTkFrame(card_res, fg_color="transparent")
    f_dim.pack(fill="x", padx=14, pady=(0, 12))
    ctk.CTkLabel(f_dim, text="Lebar:").pack(side="left", padx=(0, 4))
    self.e_w = ctk.CTkEntry(
        f_dim, textvariable=self.custom_w, width=60, state="disabled"
    )
    self.e_w.pack(side="left", padx=(0, 10))
    ctk.CTkLabel(f_dim, text="Tinggi:").pack(side="left", padx=(0, 4))
    self.e_h = ctk.CTkEntry(
        f_dim, textvariable=self.custom_h, width=60, state="disabled"
    )
    self.e_h.pack(side="left", padx=(0, 10))
    ctk.CTkLabel(
        f_dim, text="*(Aspect ratio terjaga)", text_color="#64748b"
    ).pack(side="left")

    self.lbl_status = ctk.CTkLabel(
        self, text="Siap.", text_color="#94a3b8", font=ctk.CTkFont(size=12)
    )
    self.lbl_status.pack(anchor="w", padx=14, pady=(2, 2))

    self.progress = ctk.CTkProgressBar(self, height=8, corner_radius=4)
    self.progress.pack(fill="x", padx=10, pady=(0, 12))
    self.progress.set(0)

    f_btn = ctk.CTkFrame(self, fg_color="transparent")
    f_btn.pack(fill="x", padx=10, pady=(0, 10))
    self.btn_run = ctk.CTkButton(
        f_btn,
        text="Mulai Resize Video",
        height=40,
        font=ctk.CTkFont(size=13, weight="bold"),
        state="disabled",
        cursor="hand2",
        command=self.start_batch,
    )
    self.btn_run.pack(side="left", fill="x", expand=True, padx=(0, 6))
    self.btn_stop = ctk.CTkButton(
        f_btn,
        text="Hentikan",
        height=40,
        font=ctk.CTkFont(size=13, weight="bold"),
        fg_color="#dc2626",
        hover_color="#b91c1c",
        state="disabled",
        cursor="hand2",
        command=self.stop_batch,
    )
    self.btn_stop.pack(side="right", padx=(6, 0))

  def on_drop(self, event):
    paths = parse_dnd_paths(event.data)
    valid = []
    for item in paths:
      if os.path.isdir(item):
        for f in os.listdir(item):
          if (
              os.path.splitext(f)[1].lower().replace(".", "") in self.video_exts
          ):
            valid.append(os.path.join(item, f))
      elif os.path.isfile(item):
        if (
            os.path.splitext(item)[1].lower().replace(".", "")
            in self.video_exts
        ):
          valid.append(item)

    if valid:
      self.valid_files = valid
      self.source_label_text.set(f"{len(valid)} video terpilih.")
      if not self.folder_output.get():
        self.folder_output.set(os.path.dirname(valid[0]))
      self.lbl_detect.configure(
          text=f"Terdeteksi {len(valid)} video siap di-encode.",
          text_color="#38bdf8",
      )
      self.check_state()

  def on_preset_change(self, val):
    st = "normal" if "Custom" in val else "disabled"
    self.e_w.configure(state=st)
    self.e_h.configure(state=st)

  def select_files(self):
    files = filedialog.askopenfilenames(
        title="Pilih File Video",
        filetypes=[
            (
                "Video Files",
                "*.mp4 *.mkv *.mov *.avi *.webm *.flv *.wmv *.m4v *.ts *.3gp *.mpg",
            ),
            ("All Files", "*.*"),
        ],
    )
    if files:
      self.valid_files = list(files)
      self.source_label_text.set(f"{len(files)} file video terpilih.")
      if not self.folder_output.get():
        self.folder_output.set(os.path.dirname(files[0]))
      self.lbl_detect.configure(
          text=f"Terpilih {len(files)} video.", text_color="#38bdf8"
      )
      self.check_state()

  def select_out(self):
    d = filedialog.askdirectory(title="Pilih Folder Hasil")
    if d:
      self.folder_output.set(d)
      self.check_state()

  def check_state(self):
    if (
        self.valid_files
        and self.folder_output.get().strip()
        and self.ffmpeg_available
    ):
      self.btn_run.configure(state="normal")
      self.lbl_status.configure(
          text=f"Siap memproses {len(self.valid_files)} video.",
          text_color="#f8fafc",
      )
    else:
      self.btn_run.configure(state="disabled")

  def stop_batch(self):
    self.cancel_event.set()
    self.btn_stop.configure(state="disabled")
    if self.current_process and self.current_process.poll() is None:
      self.current_process.terminate()

  def start_batch(self):
    self.cancel_event.clear()
    self.btn_run.configure(state="disabled")
    self.btn_stop.configure(state="normal")
    threading.Thread(target=self._process, daemon=True).start()

  def _process(self):
    val = self.preset_var.get()
    w, h = (
        (1920, 1080)
        if "1080p" in val
        else (1280, 720)
        if "720p" in val
        else (854, 480)
        if "480p" in val
        else (640, 360)
        if "360p" in val
        else (int(self.custom_w.get()), int(self.custom_h.get()))
    )

    scale_f = f"scale={w}:{h}:force_original_aspect_ratio=decrease,pad=ceil(iw/2)*2:ceil(ih/2)*2"
    out_dir = self.folder_output.get()
    total = len(self.valid_files)
    ok_count = 0
    failed_items = []
    cancelled = False

    for idx, src_p in enumerate(self.valid_files, start=1):
      if self.cancel_event.is_set():
        cancelled = True
        break
      fname = os.path.basename(src_p)
      self.lbl_status.configure(text=f"Encoding ({idx}/{total}): {fname}")
      dest_p = os.path.join(out_dir, f"{os.path.splitext(fname)[0]}_resized.mp4")

      cmd = [
          get_ffmpeg_binary(),
          "-y",
          "-i",
          src_p,
          "-vf",
          scale_f,
          "-c:v",
          "libx264",
          "-preset",
          "ultrafast",
          "-crf",
          "24",
          "-pix_fmt",
          "yuv420p",
          "-c:a",
          "aac",
          "-b:a",
          "128k",
          dest_p,
      ]

      try:
        self.current_process = subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            creationflags=(
                subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
            ),
        )
        _, err = self.current_process.communicate()
        if self.current_process.returncode == 0:
          ok_count += 1
        elif self.cancel_event.is_set():
          if os.path.exists(dest_p):
            os.remove(dest_p)
          cancelled = True
          break
        else:
          failed_items.append((fname, err.decode(errors="ignore")[-200:]))
      except Exception as ex:
        failed_items.append((fname, str(ex)))

      self.progress.set(idx / total)

    self.btn_run.configure(state="normal")
    self.btn_stop.configure(state="disabled")

    if failed_items:
      log_path = os.path.join(out_dir, "video_error_log.txt")
      with open(log_path, "w", encoding="utf-8") as lf:
        lf.write("LAPORAN VIDEO GAGAL DI-ENCODE:\n" + "=" * 40 + "\n")
        for fn, err in failed_items:
          lf.write(f"- {fn}: {err}\n")

    msg = f"Selesai! {ok_count} dari {total} video berhasil di-resize."
    if failed_items:
      msg += f"\n\n{len(failed_items)} video gagal (tersimpan di video_error_log.txt)."
    msg += "\n\nBuka folder hasil sekarang?"

    if messagebox.askyesno("Status Video", msg):
      open_folder_in_explorer(out_dir)


# =============================================================
# TAB 3: MODUL AUDIO UTILITY (DND SUPPORTED)
# =============================================================
class AudioTab(ctk.CTkFrame):

  def __init__(self, parent, ffmpeg_available):
    super().__init__(parent, fg_color="transparent")
    self.ffmpeg_available = ffmpeg_available
    self.file_input = ctk.StringVar(value="Seret file media ke sini...")
    self.folder_output = ctk.StringVar()
    self.target_format = ctk.StringVar(value="MP3")
    self.target_bitrate = ctk.StringVar(value="192 kbps (Standard)")

    self.current_process = None
    self.setup_ui()

  def setup_ui(self):
    if not self.ffmpeg_available:
      f_warn = ctk.CTkFrame(self, fg_color="#ef4444", corner_radius=8)
      f_warn.pack(fill="x", padx=10, pady=(6, 0))
      ctk.CTkLabel(
          f_warn,
          text=(
              "⚠️ Modul Audio membutuhkan FFmpeg yang valid terinstal di sistem."
          ),
          text_color="white",
          font=ctk.CTkFont(size=11, weight="bold"),
      ).pack(pady=6)

    self.drop_zone = ctk.CTkFrame(
        self, corner_radius=12, border_width=2, border_color="#334155"
    )
    self.drop_zone.pack(fill="x", padx=10, pady=(10, 8))

    ctk.CTkLabel(
        self.drop_zone,
        text="SUMBER MEDIA (DRAG & DROP VIDEO / AUDIO)",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=14, pady=(10, 2))

    f_in = ctk.CTkFrame(self.drop_zone, fg_color="transparent")
    f_in.pack(fill="x", padx=14, pady=(0, 6))
    self.entry_in = ctk.CTkEntry(
        f_in, textvariable=self.file_input, state="readonly", height=34
    )
    self.entry_in.pack(side="left", fill="x", expand=True, padx=(0, 8))
    ctk.CTkButton(
        f_in,
        text="Pilih",
        width=75,
        height=34,
        command=self.select_file,
        cursor="hand2",
    ).pack(side="right")

    ctk.CTkLabel(
        self.drop_zone,
        text="FOLDER SIMPAN AUDIO",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=14, pady=(4, 2))
    f_out = ctk.CTkFrame(self.drop_zone, fg_color="transparent")
    f_out.pack(fill="x", padx=14, pady=(0, 12))
    ctk.CTkEntry(
        f_out, textvariable=self.folder_output, state="readonly", height=34
    ).pack(side="left", fill="x", expand=True, padx=(0, 8))
    ctk.CTkButton(
        f_out,
        text="Pilih",
        width=75,
        height=34,
        command=self.select_dir,
        cursor="hand2",
    ).pack(side="right")

    self.drop_zone.drop_target_register(DND_FILES)
    self.drop_zone.dnd_bind("<<Drop>>", self.on_drop)
    self.entry_in.drop_target_register(DND_FILES)
    self.entry_in.dnd_bind("<<Drop>>", self.on_drop)

    card_opt = ctk.CTkFrame(self, corner_radius=12)
    card_opt.pack(fill="x", padx=10, pady=(0, 12))

    f_grid = ctk.CTkFrame(card_opt, fg_color="transparent")
    f_grid.pack(fill="x", padx=14, pady=12)
    ctk.CTkLabel(
        f_grid, text="Format:", font=ctk.CTkFont(size=12, weight="bold")
    ).pack(side="left", padx=(0, 6))
    ctk.CTkOptionMenu(
        f_grid,
        variable=self.target_format,
        values=["MP3", "AAC", "M4A", "WAV", "FLAC", "OGG"],
        width=90,
    ).pack(side="left", padx=(0, 14))

    ctk.CTkLabel(
        f_grid, text="Bitrate:", font=ctk.CTkFont(size=12, weight="bold")
    ).pack(side="left", padx=(0, 6))
    ctk.CTkOptionMenu(
        f_grid,
        variable=self.target_bitrate,
        values=[
            "320 kbps (Studio)",
            "192 kbps (Standard)",
            "128 kbps (Efisien)",
            "64 kbps (Voice Note)",
        ],
        width=150,
    ).pack(side="left")

    self.lbl_status = ctk.CTkLabel(
        self,
        text="Tarik video/audio ke kotak atas.",
        text_color="#94a3b8",
        font=ctk.CTkFont(size=12),
    )
    self.lbl_status.pack(anchor="w", padx=14, pady=(2, 8))

    self.btn_run = ctk.CTkButton(
        self,
        text="Ekstrak & Konversi Audio",
        height=40,
        font=ctk.CTkFont(size=13, weight="bold"),
        state="disabled",
        cursor="hand2",
        command=self.start_extract,
    )
    self.btn_run.pack(fill="x", padx=10)

  def on_drop(self, event):
    paths = parse_dnd_paths(event.data)
    if paths and os.path.isfile(paths[0]):
      p = paths[0]
      self.file_input.set(p)
      if not self.folder_output.get():
        self.folder_output.set(os.path.dirname(p))
      if self.ffmpeg_available:
        self.btn_run.configure(state="normal")
      self.lbl_status.configure(
          text=f"Siap diekstrak: {os.path.basename(p)}", text_color="#38bdf8"
      )

  def select_file(self):
    p = filedialog.askopenfilename(
        title="Pilih File Media",
        filetypes=[
            (
                "Media Files",
                "*.mp4 *.mkv *.mov *.avi *.webm *.flv *.ts *.mp3 *.wav *.m4a",
            ),
            ("All Files", "*.*"),
        ],
    )
    if p:
      self.file_input.set(p)
      if not self.folder_output.get():
        self.folder_output.set(os.path.dirname(p))
      if self.ffmpeg_available:
        self.btn_run.configure(state="normal")
      self.lbl_status.configure(
          text=f"Siap diekstrak: {os.path.basename(p)}", text_color="#38bdf8"
      )

  def select_dir(self):
    d = filedialog.askdirectory(title="Pilih Folder Simpan")
    if d:
      self.folder_output.set(d)

  def start_extract(self):
    self.btn_run.configure(state="disabled")
    self.lbl_status.configure(
        text="Sedang mengekstrak audio...", text_color="#38bdf8"
    )
    threading.Thread(target=self._process, daemon=True).start()

  def _process(self):
    src = self.file_input.get()
    out = self.folder_output.get()
    fmt = self.target_format.get().lower()
    br_str = self.target_bitrate.get()
    bitrate = (
        "320k"
        if "320" in br_str
        else "192k"
        if "192" in br_str
        else "128k"
        if "128" in br_str
        else "64k"
    )

    base = os.path.splitext(os.path.basename(src))[0]
    dest = os.path.join(out, f"{base}_audio.{fmt}")

    cmd = [get_ffmpeg_binary(), "-y", "-i", src, "-vn"]
    if fmt == "mp3":
      cmd += ["-c:a", "libmp3lame", "-b:a", bitrate]
    elif fmt in ["aac", "m4a"]:
      cmd += ["-c:a", "aac", "-b:a", bitrate]
    elif fmt == "flac":
      cmd += ["-c:a", "flac"]
    elif fmt == "wav":
      cmd += ["-c:a", "pcm_s16le"]
    else:
      cmd += ["-b:a", bitrate]
    cmd.append(dest)

    try:
      self.current_process = subprocess.Popen(
          cmd,
          stdout=subprocess.PIPE,
          stderr=subprocess.PIPE,
          creationflags=(
              subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
          ),
      )
      self.current_process.wait()
      if self.current_process.returncode == 0 and os.path.exists(dest):
        self.lbl_status.configure(
            text="Ekstraksi Berhasil!", text_color="#22c55e"
        )
        if messagebox.askyesno(
            "Sukses",
            f"Audio berhasil diekstrak ke:\n{dest}\n\nBuka folder penyimpanan?",
        ):
          open_folder_in_explorer(out)
      else:
        self.lbl_status.configure(
            text="Gagal mengekstrak.", text_color="#ef4444"
        )
        messagebox.showerror("Error", "FFmpeg gagal memproses file.")
    except Exception as ex:
      self.lbl_status.configure(text="Error sistem.", text_color="#ef4444")
      messagebox.showerror("Error", str(ex))
    finally:
      self.btn_run.configure(state="normal")


# =============================================================
# TAB 4: MODUL DOKUMEN & PDF TOOLS (DND SUPPORTED)
# =============================================================
class DocumentTab(ctk.CTkFrame):

  def __init__(self, parent):
    super().__init__(parent, fg_color="transparent")
    self.file_input = ctk.StringVar(
        value="Seret file dokumen (DOCX, PDF, XLSX, dll) ke sini..."
    )
    self.folder_output = ctk.StringVar()
    self.target_format = ctk.StringVar()
    self.pdf_mode_var = ctk.StringVar(value="Visual (Gambar per Halaman)")

    self.merge_files = []
    self.split_file = ctk.StringVar(value="Seret file PDF ke sini...")
    self.split_range = ctk.StringVar(value="1-3")

    self.format_matrix = {
        "docx": ["PDF", "TXT", "HTML", "RTF", "DOC"],
        "doc": ["PDF", "DOCX", "TXT", "RTF"],
        "rtf": ["PDF", "DOCX", "TXT"],
        "txt": ["PDF", "DOCX", "HTML"],
        "pdf": ["DOCX", "TXT"],
        "xlsx": ["PDF", "CSV", "HTML", "TXT"],
        "xls": ["PDF", "XLSX", "CSV"],
        "csv": ["XLSX", "PDF"],
        "pptx": ["PDF"],
        "ppt": ["PDF", "PPTX"],
    }
    self.setup_ui()

  def setup_ui(self):
    self.sub_tabs = ctk.CTkTabview(
        self,
        corner_radius=10,
        segmented_button_selected_color="#0284c7",
        height=400,
    )
    self.sub_tabs.pack(fill="both", expand=True, padx=4, pady=4)

    tab_conv = self.sub_tabs.add("🔄 Konversi Format")
    tab_merge = self.sub_tabs.add("📑 Gabung PDF (Merge)")
    tab_split = self.sub_tabs.add("✂️ Pisah PDF (Split)")

    # --- SUBTAB 1: KONVERSI DOKUMEN ---
    self.drop_doc = ctk.CTkFrame(
        tab_conv, corner_radius=10, border_width=2, border_color="#334155"
    )
    self.drop_doc.pack(fill="x", padx=6, pady=(6, 8))

    ctk.CTkLabel(
        self.drop_doc,
        text="FILE DOKUMEN (DRAG & DROP)",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=12, pady=(8, 2))
    f_in = ctk.CTkFrame(self.drop_doc, fg_color="transparent")
    f_in.pack(fill="x", padx=12, pady=(0, 6))
    self.entry_doc = ctk.CTkEntry(
        f_in, textvariable=self.file_input, state="readonly", height=32
    )
    self.entry_doc.pack(side="left", fill="x", expand=True, padx=(0, 8))
    ctk.CTkButton(
        f_in,
        text="Pilih",
        width=75,
        height=32,
        command=self.select_file,
        cursor="hand2",
    ).pack(side="right")

    ctk.CTkLabel(
        self.drop_doc,
        text="FOLDER SIMPAN HASIL",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=12, pady=(0, 2))
    f_out = ctk.CTkFrame(self.drop_doc, fg_color="transparent")
    f_out.pack(fill="x", padx=12, pady=(0, 10))
    ctk.CTkEntry(
        f_out, textvariable=self.folder_output, state="readonly", height=32
    ).pack(side="left", fill="x", expand=True, padx=(0, 8))
    ctk.CTkButton(
        f_out,
        text="Pilih",
        width=75,
        height=32,
        command=self.select_dir,
        cursor="hand2",
    ).pack(side="right")

    self.drop_doc.drop_target_register(DND_FILES)
    self.drop_doc.dnd_bind("<<Drop>>", self.on_drop_doc)
    self.entry_doc.drop_target_register(DND_FILES)
    self.entry_doc.dnd_bind("<<Drop>>", self.on_drop_doc)

    card_opt = ctk.CTkFrame(tab_conv, corner_radius=10)
    card_opt.pack(fill="x", padx=6, pady=(0, 8))
    ctk.CTkLabel(
        card_opt,
        text="FORMAT TARGET",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=12, pady=(8, 2))
    self.cb_fmt = ctk.CTkOptionMenu(
        card_opt,
        variable=self.target_format,
        values=["Pilih Dokumen Dulu"],
        state="disabled",
        height=32,
        command=self.on_target_format_change,
    )
    self.cb_fmt.pack(fill="x", padx=12, pady=(0, 4))

    f_pdf_mode = ctk.CTkFrame(card_opt, fg_color="transparent")
    f_pdf_mode.pack(fill="x", padx=12, pady=(0, 10))
    ctk.CTkLabel(
        f_pdf_mode,
        text="Mode PDF→DOCX:",
        font=ctk.CTkFont(size=12, weight="bold"),
        text_color="#94a3b8",
    ).pack(side="left", padx=(0, 8))
    self.cb_pdf_mode = ctk.CTkOptionMenu(
        f_pdf_mode,
        variable=self.pdf_mode_var,
        values=["Visual (Gambar per Halaman)", "Teks (Editable)"],
        state="disabled",
        width=240,
        height=32,
    )
    self.cb_pdf_mode.pack(side="left")

    self.lbl_conv_status = ctk.CTkLabel(
        tab_conv,
        text="Tarik dokumen atau klik Pilih File.",
        text_color="#94a3b8",
        font=ctk.CTkFont(size=11),
    )
    self.lbl_conv_status.pack(anchor="w", padx=12, pady=(0, 6))

    self.btn_conv = ctk.CTkButton(
        tab_conv,
        text="Mulai Konversi Dokumen",
        height=38,
        font=ctk.CTkFont(size=12, weight="bold"),
        state="disabled",
        cursor="hand2",
        command=self.start_convert,
    )
    self.btn_conv.pack(fill="x", padx=6)

    # --- SUBTAB 2: MERGE PDF ---
    self.drop_merge = ctk.CTkFrame(
        tab_merge, corner_radius=10, border_width=2, border_color="#334155"
    )
    self.drop_merge.pack(fill="both", expand=True, padx=6, pady=6)

    ctk.CTkLabel(
        self.drop_merge,
        text="GABUNGKAN FILE PDF (DRAG & DROP ATAU PILIH)",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=12, pady=(8, 4))
    ctk.CTkButton(
        self.drop_merge,
        text="+ Pilih File PDF Sekaligus",
        height=32,
        command=self.select_merge_files,
        cursor="hand2",
    ).pack(fill="x", padx=12, pady=(0, 6))

    self.lbl_merge_info = ctk.CTkLabel(
        self.drop_merge,
        text="Seret banyak file PDF langsung ke kotak ini.",
        text_color="#64748b",
        font=ctk.CTkFont(size=11),
    )
    self.lbl_merge_info.pack(anchor="w", padx=12, pady=(0, 8))

    self.btn_merge = ctk.CTkButton(
        self.drop_merge,
        text="Gabungkan Menjadi 1 File PDF",
        height=38,
        fg_color="#0284c7",
        font=ctk.CTkFont(size=12, weight="bold"),
        state="disabled",
        cursor="hand2",
        command=self.process_merge_pdf,
    )
    self.btn_merge.pack(fill="x", padx=12, pady=(0, 10))

    self.drop_merge.drop_target_register(DND_FILES)
    self.drop_merge.dnd_bind("<<Drop>>", self.on_drop_merge)

    # --- SUBTAB 3: SPLIT PDF ---
    self.drop_split = ctk.CTkFrame(
        tab_split, corner_radius=10, border_width=2, border_color="#334155"
    )
    self.drop_split.pack(fill="both", expand=True, padx=6, pady=6)

    ctk.CTkLabel(
        self.drop_split,
        text="PISAH HALAMAN PDF (DRAG & DROP ATAU PILIH)",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=12, pady=(8, 4))
    f_sin = ctk.CTkFrame(self.drop_split, fg_color="transparent")
    f_sin.pack(fill="x", padx=12, pady=(0, 6))
    self.entry_split = ctk.CTkEntry(
        f_sin, textvariable=self.split_file, state="readonly", height=32
    )
    self.entry_split.pack(side="left", fill="x", expand=True, padx=(0, 8))
    ctk.CTkButton(
        f_sin,
        text="Pilih",
        width=75,
        height=32,
        command=self.select_split_file,
        cursor="hand2",
    ).pack(side="right")

    f_range = ctk.CTkFrame(self.drop_split, fg_color="transparent")
    f_range.pack(fill="x", padx=12, pady=(4, 10))
    ctk.CTkLabel(
        f_range,
        text="Rentang Halaman (contoh: 1-3 atau 1,4,5):",
        font=ctk.CTkFont(size=11),
    ).pack(side="left", padx=(0, 8))
    ctk.CTkEntry(
        f_range, textvariable=self.split_range, width=120, height=32
    ).pack(side="left")

    self.btn_split = ctk.CTkButton(
        self.drop_split,
        text="Ekstrak & Simpan PDF Baru",
        height=38,
        fg_color="#0284c7",
        font=ctk.CTkFont(size=12, weight="bold"),
        state="disabled",
        cursor="hand2",
        command=self.process_split_pdf,
    )
    self.btn_split.pack(fill="x", padx=12, pady=(0, 10))

    self.drop_split.drop_target_register(DND_FILES)
    self.drop_split.dnd_bind("<<Drop>>", self.on_drop_split)
    self.entry_split.drop_target_register(DND_FILES)
    self.entry_split.dnd_bind("<<Drop>>", self.on_drop_split)

  def on_drop_doc(self, event):
    paths = parse_dnd_paths(event.data)
    if paths and os.path.isfile(paths[0]):
      self._load_document(paths[0])

  def _load_document(self, p):
    ext = os.path.splitext(p)[1].lower().replace(".", "")
    if ext in self.format_matrix:
      self.file_input.set(p)
      if not self.folder_output.get():
        self.folder_output.set(os.path.dirname(p))
      targets = self.format_matrix[ext]
      self.cb_fmt.configure(state="normal", values=targets)
      self.target_format.set(targets[0])
      self.btn_conv.configure(state="normal")
      self.lbl_conv_status.configure(
          text=f"File terpilih: .{ext.upper()}", text_color="#38bdf8"
      )
      self._update_pdf_mode_state(ext)
    else:
      messagebox.showerror("Error", f"Format .{ext} belum didukung.")

  def on_target_format_change(self, val):
    ext = os.path.splitext(self.file_input.get())[1].lower().replace(".", "")
    self._update_pdf_mode_state(ext)

  def _update_pdf_mode_state(self, src_ext):
    if src_ext == "pdf" and self.target_format.get().lower() == "docx":
      self.cb_pdf_mode.configure(state="normal")
    else:
      self.cb_pdf_mode.configure(state="disabled")

  def on_drop_merge(self, event):
    paths = parse_dnd_paths(event.data)
    pdf_files = [
        p for p in paths if os.path.isfile(p) and p.lower().endswith(".pdf")
    ]
    if len(pdf_files) >= 2:
      self.merge_files = pdf_files
      self.lbl_merge_info.configure(
          text=f"Terdeteksi {len(pdf_files)} file PDF untuk digabung.",
          text_color="#38bdf8",
      )
      self.btn_merge.configure(state="normal")
    else:
      messagebox.showwarning(
          "Peringatan", "Seret minimal 2 file PDF untuk digabungkan."
      )

  def on_drop_split(self, event):
    paths = parse_dnd_paths(event.data)
    if (
        paths
        and os.path.isfile(paths[0])
        and paths[0].lower().endswith(".pdf")
    ):
      self.split_file.set(paths[0])
      self.btn_split.configure(state="normal")

  def select_file(self):
    exts = " ".join([f"*.{k}" for k in self.format_matrix.keys()])
    p = filedialog.askopenfilename(
        title="Pilih Dokumen",
        filetypes=[("Dokumen Didukung", exts), ("Semua File", "*.*")],
    )
    if p:
      self._load_document(p)

  def select_dir(self):
    d = filedialog.askdirectory(title="Pilih Folder Simpan")
    if d:
      self.folder_output.set(d)

  def start_convert(self):
    self.btn_conv.configure(state="disabled")
    self.lbl_conv_status.configure(
        text="Sedang memproses konversi...", text_color="#38bdf8"
    )
    threading.Thread(target=self._process_conv, daemon=True).start()

  def _process_conv(self):
    src = self.file_input.get()
    out = self.folder_output.get()
    t_fmt = self.target_format.get().lower()
    src_ext = os.path.splitext(src)[1].lower().replace(".", "")
    dest = os.path.join(
        out, f"{os.path.splitext(os.path.basename(src))[0]}.{t_fmt}"
    )

    success = False
    try:
      # Jalur 1: PDF ke DOCX atau TXT (Pustaka murni)
      if src_ext == "pdf":
        if t_fmt == "docx":
          if self.pdf_mode_var.get() == "Visual (Gambar per Halaman)":
            self._convert_pdf_visual(src, dest)
            success = os.path.exists(dest)
          elif HAS_PDF2DOCX:
            cv = PdfToDocxConverter(src)
            cv.convert(dest, start=0, end=None)
            cv.close()
            success = True
          else:
            messagebox.showerror(
                "Pustaka Hilang",
                "Library pdf2docx tidak tersedia. Jalankan: pip install pdf2docx",
            )
        elif t_fmt == "txt" and HAS_PYPDF:
          reader = PdfReader(src)
          txt = "\n\n".join(
              [p.extract_text() for p in reader.pages if p.extract_text()]
          )
          with open(dest, "w", encoding="utf-8") as f:
            f.write(txt)
          success = True

      # Jalur 2: Dokumen Office di Windows via COM
      elif HAS_WIN32COM and platform.system() == "Windows":
        pythoncom.CoInitialize()
        if src_ext in ["docx", "doc", "rtf", "txt"]:
          w = win32com.client.DispatchEx("Word.Application")
          w.Visible = False
          d = w.Documents.Open(os.path.abspath(src))
          fmap = {
              "pdf": 17, "docx": 16, "txt": 2, "rtf": 6, "html": 8, "doc": 0,
          }
          d.SaveAs(os.path.abspath(dest), FileFormat=fmap.get(t_fmt, 17))
          d.Close()
          w.Quit()
          success = True
        elif src_ext in ["xlsx", "xls", "csv"]:
          xl = win32com.client.DispatchEx("Excel.Application")
          xl.Visible = False
          wb = xl.Workbooks.Open(os.path.abspath(src))
          if t_fmt == "pdf":
            wb.ExportAsFixedFormat(0, os.path.abspath(dest))
          elif t_fmt == "csv":
            wb.SaveAs(os.path.abspath(dest), FileFormat=6)
          elif t_fmt == "xlsx":
            wb.SaveAs(os.path.abspath(dest), FileFormat=51)
          wb.Close(False)
          xl.Quit()
          success = True
        elif src_ext in ["pptx", "ppt"]:
          ppt = win32com.client.DispatchEx("PowerPoint.Application")
          prs = ppt.Presentations.Open(os.path.abspath(src), WithWindow=False)
          if t_fmt == "pdf":
            prs.SaveAs(os.path.abspath(dest), 32)
          prs.Close()
          ppt.Quit()
          success = True
        pythoncom.CoUninitialize()

      # Jalur 3: Dokumen Office di Linux/macOS via LibreOffice Headless
      else:
        soffice_bin = get_libreoffice_command()
        if soffice_bin:
          cmd = [
              soffice_bin,
              "--headless",
              "--convert-to", t_fmt,
              src,
              "--outdir", out
          ]
          proc = subprocess.run(
              cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE
          )
          success = (proc.returncode == 0)

      if success and os.path.exists(dest):
        self.lbl_conv_status.configure(
            text="Konversi Berhasil!", text_color="#22c55e"
        )
        if messagebox.askyesno(
            "Sukses", f"Dokumen disimpan di:\n{dest}\n\nBuka folder hasil?"
        ):
          open_folder_in_explorer(out)
      else:
        self.lbl_conv_status.configure(
            text="Konversi Gagal.", text_color="#ef4444"
        )
        messagebox.showerror(
            "Gagal",
            "Gagal memproses dokumen.\nPastikan Microsoft Office (Windows) atau LibreOffice (Linux/Mac) terpasang."
        )
    except Exception as ex:
      self.lbl_conv_status.configure(text="Error konversi.", text_color="#ef4444")
      messagebox.showerror("Error", str(ex))
    finally:
      self.btn_conv.configure(state="normal")

  def _convert_pdf_visual(self, src, dest):
    """Render tiap halaman PDF menjadi gambar dan sisipkan ke DOCX (hasil persis)."""
    if not HAS_FITZ or not HAS_DOCX:
      raise RuntimeError(
          "Mode Visual membutuhkan PyMuPDF & python-docx.\nBila sudah ada, gunakan Mode Teks (Editable)."
      )
    zoom = 2.0  # ~144 DPI: tajam & muat di halaman A4 tanpa perlu margin 0 rapi
    dpi_mm = 72.0 * zoom

    pdf_doc = fitz.open(src)
    if pdf_doc.page_count == 0:
      pdf_doc.close()
      raise RuntimeError("PDF tidak memiliki halaman.")

    docx_doc = DocxDocument()
    sec0 = docx_doc.sections[0]
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
      setattr(sec0, m, Mm(0))

    for i, page in enumerate(pdf_doc):
      pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=True)
      page_w_mm = pix.width / dpi_mm * 25.4
      page_h_mm = pix.height / dpi_mm * 25.4

      if i == 0:
        sec = sec0
      else:
        sec = docx_doc.add_section(WD_SECTION.NEW_PAGE)
      sec.page_width = Mm(page_w_mm)
      sec.page_height = Mm(page_h_mm)
      for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Mm(0))
      run = docx_doc.add_paragraph().add_run()
      run.add_picture(BytesIO(pix.tobytes("png")), width=Mm(page_w_mm))

    pdf_doc.close()
    docx_doc.save(dest)

  def select_merge_files(self):
    files = filedialog.askopenfilenames(
        title="Pilih File PDF", filetypes=[("PDF Files", "*.pdf")]
    )
    if files and len(files) >= 2:
      self.merge_files = list(files)
      self.lbl_merge_info.configure(
          text=f"Terpilih {len(self.merge_files)} PDF untuk digabung.",
          text_color="#38bdf8",
      )
      self.btn_merge.configure(state="normal")

  def process_merge_pdf(self):
    dest = filedialog.asksaveasfilename(
        title="Simpan PDF Gabungan",
        defaultextension=".pdf",
        filetypes=[("PDF Files", "*.pdf")],
    )
    if not dest:
      return
    try:
      writer = PdfWriter()
      for f in self.merge_files:
        reader = PdfReader(f)
        for page in reader.pages:
          writer.add_page(page)
      with open(dest, "wb") as out_f:
        writer.write(out_f)
      if messagebox.askyesno(
          "Sukses",
          f"PDF berhasil digabung ke:\n{dest}\n\nBuka folder penyimpanan?",
      ):
        open_folder_in_explorer(os.path.dirname(dest))
    except Exception as e:
      messagebox.showerror("Error", f"Gagal gabung PDF: {e}")

  def select_split_file(self):
    f = filedialog.askopenfilename(
        title="Pilih PDF", filetypes=[("PDF Files", "*.pdf")]
    )
    if f:
      self.split_file.set(f)
      self.btn_split.configure(state="normal")

  def process_split_pdf(self):
    src = self.split_file.get()
    raw_range = self.split_range.get().strip()
    dest = filedialog.asksaveasfilename(
        title="Simpan PDF Hasil Pisah",
        defaultextension=".pdf",
        filetypes=[("PDF Files", "*.pdf")],
    )
    if not dest:
      return
    try:
      reader = PdfReader(src)
      total = len(reader.pages)
      pages = set()
      for part in raw_range.split(","):
        part = part.strip()
        if "-" in part:
          s, e = part.split("-")
          for p in range(int(s), int(e) + 1):
            if 1 <= p <= total:
              pages.add(p - 1)
        else:
          p = int(part)
          if 1 <= p <= total:
            pages.add(p - 1)

      writer = PdfWriter()
      for idx in sorted(list(pages)):
        writer.add_page(reader.pages[idx])
      with open(dest, "wb") as out_f:
        writer.write(out_f)
      if messagebox.askyesno(
          "Sukses",
          f"Halaman PDF tersimpan di:\n{dest}\n\nBuka folder penyimpanan?",
      ):
        open_folder_in_explorer(os.path.dirname(dest))
    except Exception as e:
      messagebox.showerror("Error", f"Gagal pisah PDF: {e}")


# =============================================================
# TAB 5: MODUL MEDIA DOWNLOADER (YOUTUBE / TIKTOK / DLL)
# =============================================================
class DownloaderTab(ctk.CTkFrame):

  def __init__(self, parent, ffmpeg_available=True):
    super().__init__(parent, fg_color="transparent")
    self.ffmpeg_available = ffmpeg_available

    self.url_var = ctk.StringVar()
    self.folder_output = ctk.StringVar()
    self.download_type = ctk.StringVar(value="Video (MP4)")
    self.video_quality = ctk.StringVar(value="Terbaik (Auto Max)")

    self.cancel_flag = False

    self.setup_ui()

  def setup_ui(self):
    card_input = ctk.CTkFrame(self, corner_radius=12)
    card_input.pack(fill="x", padx=10, pady=(10, 8))

    ctk.CTkLabel(
        card_input,
        text="LINK / URL MEDIA (YOUTUBE, TIKTOK, IG, TWITTER, SOUNDCLOUD, DLL)",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=14, pady=(10, 2))

    f_url = ctk.CTkFrame(card_input, fg_color="transparent")
    f_url.pack(fill="x", padx=14, pady=(0, 8))

    self.entry_url = ctk.CTkEntry(
        f_url,
        textvariable=self.url_var,
        placeholder_text="Tempel link video/musik di sini (https://...)",
        height=36,
    )
    self.entry_url.pack(side="left", fill="x", expand=True, padx=(0, 8))

    ctk.CTkButton(
        f_url,
        text="Paste",
        width=70,
        height=36,
        command=self.paste_from_clipboard,
        cursor="hand2",
    ).pack(side="right")

    ctk.CTkLabel(
        card_input,
        text="FOLDER SIMPAN HASIL UNDUHAN",
        font=ctk.CTkFont(size=11, weight="bold"),
        text_color="#94a3b8",
    ).pack(anchor="w", padx=14, pady=(0, 2))

    f_out = ctk.CTkFrame(card_input, fg_color="transparent")
    f_out.pack(fill="x", padx=14, pady=(0, 14))

    self.entry_out = ctk.CTkEntry(
        f_out, textvariable=self.folder_output, state="readonly", height=34
    )
    self.entry_out.pack(side="left", fill="x", expand=True, padx=(0, 8))

    ctk.CTkButton(
        f_out,
        text="Pilih",
        width=70,
        height=34,
        command=self.select_out_dir,
        cursor="hand2",
    ).pack(side="right")

    default_download_dir = os.path.join(os.path.expanduser("~"), "Downloads")
    self.folder_output.set(default_download_dir)

    card_opt = ctk.CTkFrame(self, corner_radius=12)
    card_opt.pack(fill="x", padx=10, pady=(0, 10))

    f_grid = ctk.CTkFrame(card_opt, fg_color="transparent")
    f_grid.pack(fill="x", padx=14, pady=12)

    ctk.CTkLabel(
        f_grid,
        text="Tipe Unduhan:",
        font=ctk.CTkFont(size=12, weight="bold"),
    ).pack(side="left", padx=(0, 6))

    self.cb_type = ctk.CTkOptionMenu(
        f_grid,
        variable=self.download_type,
        values=["Video (MP4)", "Musik/Audio Saja (MP3)"],
        command=self.on_type_change,
        width=170,
    )
    self.cb_type.pack(side="left", padx=(0, 16))

    self.lbl_quality = ctk.CTkLabel(
        f_grid,
        text="Kualitas:",
        font=ctk.CTkFont(size=12, weight="bold"),
    )
    self.lbl_quality.pack(side="left", padx=(0, 6))

    self.cb_quality = ctk.CTkOptionMenu(
        f_grid,
        variable=self.video_quality,
        values=[
            "Terbaik (Auto Max)",
            "1080p (FHD)",
            "720p (HD)",
            "480p (SD)",
            "360p",
        ],
        width=150,
    )
    self.cb_quality.pack(side="left")

    self.lbl_status = ctk.CTkLabel(
        self,
        text="Siap mengunduh. Tempel tautan URL di atas.",
        text_color="#94a3b8",
        font=ctk.CTkFont(size=12),
    )
    self.lbl_status.pack(anchor="w", padx=14, pady=(2, 2))

    self.progress = ctk.CTkProgressBar(self, height=8, corner_radius=4)
    self.progress.pack(fill="x", padx=10, pady=(0, 12))
    self.progress.set(0)

    f_btn = ctk.CTkFrame(self, fg_color="transparent")
    f_btn.pack(fill="x", padx=10, pady=(0, 10))

    self.btn_download = ctk.CTkButton(
        f_btn,
        text="Mulai Download",
        height=40,
        font=ctk.CTkFont(size=13, weight="bold"),
        cursor="hand2",
        command=self.start_download_thread,
    )
    self.btn_download.pack(side="left", fill="x", expand=True, padx=(0, 6))

    self.btn_cancel = ctk.CTkButton(
        f_btn,
        text="Batal",
        height=40,
        font=ctk.CTkFont(size=13, weight="bold"),
        fg_color="#dc2626",
        hover_color="#b91c1c",
        state="disabled",
        cursor="hand2",
        command=self.cancel_download,
    )
    self.btn_cancel.pack(side="right", padx=(6, 0))

  def paste_from_clipboard(self):
    try:
      text = self.clipboard_get().strip()
      if text.startswith("http://") or text.startswith("https://"):
        self.url_var.set(text)
        self.lbl_status.configure(
            text="Tautan siap diunduh.", text_color="#38bdf8"
        )
      else:
        messagebox.showwarning(
            "Clipboard", "Isi clipboard bukan format URL yang valid."
        )
    except Exception:
      pass

  def select_out_dir(self):
    d = filedialog.askdirectory(title="Pilih Folder Penyimpanan")
    if d:
      self.folder_output.set(d)

  def on_type_change(self, val):
    if "Audio" in val:
      self.cb_quality.configure(state="disabled")
      self.lbl_quality.configure(text_color="#64748b")
    else:
      self.cb_quality.configure(state="normal")
      self.lbl_quality.configure(text_color="#f8fafc")

  def cancel_download(self):
    self.cancel_flag = True
    self.lbl_status.configure(
        text="Membatalkan unduhan...", text_color="#ef4444"
    )
    self.btn_cancel.configure(state="disabled")

  def start_download_thread(self):
    if not HAS_YTDLP:
      messagebox.showerror(
          "Pustaka Hilang",
          "Library yt-dlp belum terpasang. Jalankan: pip install yt-dlp",
      )
      return

    url = self.url_var.get().strip()
    if not url.startswith("http://") and not url.startswith("https://"):
      messagebox.showerror("URL Kosong", "Masukkan tautan URL media yang valid.")
      return

    self.cancel_flag = False
    self.progress.set(0)
    self.btn_download.configure(state="disabled")
    self.btn_cancel.configure(state="normal")
    self.lbl_status.configure(
        text="Mengambil informasi media...", text_color="#38bdf8"
    )

    threading.Thread(target=self._download_worker, daemon=True).start()

  def _progress_hook(self, d):
    if self.cancel_flag:
      raise yt_dlp.utils.DownloadCancelled("Dibatalkan pengguna.")

    if d["status"] == "downloading":
      total = d.get("total_bytes") or d.get("total_bytes_estimate") or 0
      downloaded = d.get("downloaded_bytes", 0)
      speed = d.get("speed") or 0
      eta = d.get("eta") or 0

      speed_str = (
          f"{speed / (1024 * 1024):.1f} MB/s" if speed else "Menghitung..."
      )

      if total > 0:
        ratio = downloaded / total
        percent = ratio * 100
        self.progress.set(ratio)
        self.lbl_status.configure(
            text=f"Mengunduh: {percent:.1f}% ({speed_str}) | Sisa: {eta} dtk",
            text_color="#38bdf8",
        )
      else:
        self.lbl_status.configure(
            text=f"Mengunduh data... ({speed_str})", text_color="#38bdf8"
        )
    elif d["status"] == "finished":
      self.progress.set(1.0)
      self.lbl_status.configure(
          text="Konversi/Muxing file selesai...", text_color="#22c55e"
      )

  def _download_worker(self):
    url = self.url_var.get().strip()
    out_dir = self.folder_output.get()
    is_audio_only = "Audio" in self.download_type.get()
    quality = self.video_quality.get()

    out_template = os.path.join(out_dir, "%(title)s.%(ext)s")

    ydl_opts = {
        "outtmpl": out_template,
        "progress_hooks": [self._progress_hook],
        "noplaylist": True,
        "quiet": True,
        "no_warnings": True,
    }

    if is_audio_only:
      ydl_opts.update({
          "format": "bestaudio/best",
          "postprocessors": [{
              "key": "FFmpegExtractAudio",
              "preferredcodec": "mp3",
              "preferredquality": "192",
          }],
      })
    else:
      if "1080p" in quality:
        f_rule = "bestvideo[height<=1080]+bestaudio/best[height<=1080]/best"
      elif "720p" in quality:
        f_rule = "bestvideo[height<=720]+bestaudio/best[height<=720]/best"
      elif "480p" in quality:
        f_rule = "bestvideo[height<=480]+bestaudio/best[height<=480]/best"
      elif "360p" in quality:
        f_rule = "bestvideo[height<=360]+bestaudio/best[height<=360]/best"
      else:
        f_rule = "bestvideo+bestaudio/best"

      ydl_opts.update({
          "format": f_rule,
          "merge_output_format": "mp4",
      })

    success = False
    cancelled = False
    err_msg = ""

    try:
      with yt_dlp.YoutubeDL(ydl_opts) as ydl:
        ydl.download([url])
      success = True
    except yt_dlp.utils.DownloadCancelled:
      cancelled = True
    except Exception as ex:
      err_msg = str(ex)

    self.btn_download.configure(state="normal")
    self.btn_cancel.configure(state="disabled")

    if cancelled:
      self.lbl_status.configure(
          text="Unduhan dibatalkan.", text_color="#ef4444"
      )
    elif success:
      self.lbl_status.configure(
          text="Unduhan Sukses!", text_color="#22c55e"
      )
      if messagebox.askyesno(
          "Sukses", "Media berhasil diunduh!\n\nBuka folder penyimpanan?"
      ):
        open_folder_in_explorer(out_dir)
    else:
      self.lbl_status.configure(
          text="Unduhan Gagal.", text_color="#ef4444"
      )
      messagebox.showerror(
          "Gagal Mengunduh",
          f"Terjadi kesalahan saat mengunduh:\n{err_msg[:300]}",
      )


# =============================================================
# JENDELA UTAMA: SUITE PRODUCTION TERPADU
# =============================================================
class UltimateMediaSuiteFull(ctk.CTk, TkinterDnD.DnDWrapper):

  def __init__(self):
    super().__init__()
    self.TkdndVersion = TkinterDnD._require(self)

    self.title("Universal Media, Document & Downloader Suite Pro")
    self.geometry("670x710")
    self.resizable(False, False)

    self.ffmpeg_available = check_ffmpeg_available()

    # Header Dashboard
    f_head = ctk.CTkFrame(self, corner_radius=0, fg_color="#0f172a", height=65)
    f_head.pack(fill="x")
    f_head.pack_propagate(False)

    ctk.CTkLabel(
        f_head,
        text="⚡ Universal Media, Document & Downloader Suite",
        font=ctk.CTkFont(size=16, weight="bold"),
        text_color="#f8fafc",
    ).pack(anchor="w", padx=18, pady=(12, 0))
    ctk.CTkLabel(
        f_head,
        text="All-in-One Local Suite: Photo/RAW, Video, Audio, Document/PDF & Media Downloader",
        font=ctk.CTkFont(size=11),
        text_color="#64748b",
    ).pack(anchor="w", padx=18, pady=(0, 8))

    # Segmented Tabview
    self.tabview = ctk.CTkTabview(
        self,
        corner_radius=12,
        segmented_button_selected_color="#0284c7",
        segmented_button_selected_hover_color="#0369a1",
    )
    self.tabview.pack(fill="both", expand=True, padx=12, pady=8)

    self.tabview.add("🖼️ Foto & RAW")
    self.tabview.add("🎬 Video")
    self.tabview.add("🎵 Audio")
    self.tabview.add("📑 Dokumen & PDF")
    self.tabview.add("📥 Downloader")

    ImageTab(self.tabview.tab("🖼️ Foto & RAW")).pack(fill="both", expand=True)
    VideoTab(
        self.tabview.tab("🎬 Video"), ffmpeg_available=self.ffmpeg_available
    ).pack(fill="both", expand=True)
    AudioTab(
        self.tabview.tab("🎵 Audio"), ffmpeg_available=self.ffmpeg_available
    ).pack(fill="both", expand=True)
    DocumentTab(self.tabview.tab("📑 Dokumen & PDF")).pack(
        fill="both", expand=True
    )
    DownloaderTab(
        self.tabview.tab("📥 Downloader"), ffmpeg_available=self.ffmpeg_available
    ).pack(fill="both", expand=True)


if __name__ == "__main__":
  app = UltimateMediaSuiteFull()
  app.mainloop()